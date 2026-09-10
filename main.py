import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Executive Navy & Slate)
COLOR_BG = RGBColor(248, 250, 252)  # Very light gray/blue background
COLOR_DARK_BG = RGBColor(15, 23, 42)  # Navy Dark
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)  # Slate 800
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
COLOR_ACCENT = RGBColor(37, 99, 235)  # Royal Blue
COLOR_ACCENT_RED = RGBColor(225, 29, 72)  # Rose Red
COLOR_CARD_BG = RGBColor(255, 255, 255)  # White
COLOR_BORDER = RGBColor(226, 232, 240)  # Light Border


def set_slide_background(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, title_text, category_text="PERFORMANCE & PROCESSOS OPERACIONAIS"):
    # Header container
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_MAIN


def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


# ==========================================
# SLIDE 1: COVER (Dark Theme, High Impact)
# ==========================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide1, COLOR_DARK_BG)

# Accent line
accent_line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.8), Inches(0.08))
accent_line.fill.solid()
accent_line.fill.fore_color.rgb = COLOR_ACCENT
accent_line.line.fill.background()

tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.3), Inches(4.0))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "Diagnóstico Operacional &\nPlano de Ação"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)
p.space_after = Pt(14)

p2 = tf1.add_paragraph()
p2.text = "Otimização de Processos e Redução de Lead Time no Cadastro de Clientes"
p2.font.size = Pt(20)
p2.font.color.rgb = RGBColor(148, 163, 184)
p2.space_after = Pt(40)

p3 = tf1.add_paragraph()
p3.text = "Apresentado por: Gilcemar  |  Área de Performance Operacional"
p3.font.size = Pt(13)
p3.font.color.rgb = RGBColor(203, 213, 225)

# ==========================================
# SLIDE 2: EXECUTIVE SUMMARY (Metric Cards)
# ==========================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide2, COLOR_BG)
add_header(slide2, "Resumo Executivo: Diagnóstico Inicial")

metrics = [
    ("3,32 h", "Lead Time Médio Atual", "Apenas 14 min são de trabalho ativo no ERP TOTVS.", COLOR_TEXT_MAIN),
    ("93%", "Tempo de Fila / Espera", "Gargalo crítico: solicitações paradas aguardando triagem.", COLOR_ACCENT_RED),
    ("30%", "Taxa de Retrabalho", "9 de 30 chamados possuem erros ou falhas na origem.", COLOR_ACCENT),
    ("89%", "Aumento de Tempo", "Casos com erro saltam de 2,62h para 4,95h de atendimento.", COLOR_ACCENT_RED)
]

card_w = Inches(2.7)
card_h = Inches(4.8)
start_x = Inches(0.8)
spacing = Inches(0.3)

for i, (val, lbl, desc, val_color) in enumerate(metrics):
    x = start_x + i * (card_w + spacing)
    add_card(slide2, x, Inches(1.6), card_w, card_h)

    tb = slide2.shapes.add_textbox(x + Inches(0.2), Inches(1.8), card_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = val
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = val_color
    p.space_after = Pt(10)

    p_lbl = tf.add_paragraph()
    p_lbl.text = lbl
    p_lbl.font.size = Pt(14)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = COLOR_TEXT_MAIN
    p_lbl.space_after = Pt(14)

    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.size = Pt(12)
    p_desc.font.color.rgb = COLOR_TEXT_MUTED

# ==========================================
# SLIDE 3: AS-IS PROCESS (Flow Cards)
# ==========================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide3, COLOR_BG)
add_header(slide3, "Mapeamento do Processo Atual (As-Is)")

# Top Banner Card
banner = add_card(slide3, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.9), bg_color=COLOR_DARK_BG)
tb = slide3.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(11.3), Inches(0.6))
tf = tb.text_frame
p = tf.paragraphs[0]
p.text = "ENTRADA DESCENTRALIZADA: E-mail não padronizado + Ficha cadastral em Word anexada"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)

# Steps
steps = [
    ("Etapa 01", "Triagem Manual", "Leitura individual do e-mail e validação inicial da ficha anexada."),
    ("Etapa 02", "Checagem TOTVS", "Verificação manual de duplicidade de CNPJ no sistema ERP."),
    ("Etapa 03", "Tratativa / Vaivém", "Troca contínua de e-mails em 30% dos casos para alinhar pendências."),
    ("Etapa 04", "Digitação Final", "Input manual dos dados validados no TOTVS (apenas 14 minutos).")
]

step_w = Inches(2.7)
step_h = Inches(3.2)

for i, (num, title, desc) in enumerate(steps):
    x = start_x + i * (step_w + spacing)
    add_card(slide3, x, Inches(2.7), step_w, step_h)

    tb = slide3.shapes.add_textbox(x + Inches(0.2), Inches(2.9), step_w - Inches(0.4), step_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(6)

    p_t = tf.add_paragraph()
    p_t.text = title
    p_t.font.size = Pt(16)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_MAIN
    p_t.space_after = Pt(10)

    p_d = tf.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(12)
    p_d.font.color.rgb = COLOR_TEXT_MUTED

# Bottom Highlight
add_card(slide3, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.8), bg_color=RGBColor(254, 242, 242),
         border_color=RGBColor(254, 202, 202))
tb_p = slide3.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.5))
p = tb_p.text_frame.paragraphs[0]
p.text = "PONTO CRÍTICO: Ausência de campos obrigatórios na origem resulta em trocas excessivas de e-mails e acúmulo de fila."
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT_RED

# ==========================================
# SLIDE 4: QUANTITATIVE DATA (Table + Chart)
# ==========================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide4, COLOR_BG)
add_header(slide4, "Análise Quantitativa da Performance")

# Left Side Table Card
add_card(slide4, Inches(0.8), Inches(1.5), Inches(6.0), Inches(4.2))
table_shape = slide4.shapes.add_table(5, 3, Inches(1.0), Inches(1.7), Inches(5.6), Inches(3.8))
table = table_shape.table

t_data = [
    ["Indicador", "Sem Pendência", "Com Pendência"],
    ["Volume / % Total", "21 casos (70%)", "9 casos (30%)"],
    ["Tempo Execução ERP", "14,14 min", "14,22 min"],
    ["Tempo de Fila / Espera", "2,38 horas", "4,71 horas"],
    ["Lead Time Total", "2,62 horas", "4,95 horas"]
]

for r_idx, row in enumerate(t_data):
    for c_idx, val in enumerate(row):
        cell = table.cell(r_idx, c_idx)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        if r_idx == 0:
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_DARK_BG
        else:
            if c_idx == 2 and r_idx >= 3:
                p.font.bold = True
                p.font.color.rgb = COLOR_ACCENT_RED

# Right Side Chart Embed
if os.path.exists('grafico_1_impacto_pendencias.png'):
    add_card(slide4, Inches(7.1), Inches(1.5), Inches(5.4), Inches(4.2))
    slide4.shapes.add_picture('grafico_1_impacto_pendencias.png', Inches(7.2), Inches(1.6), Inches(5.2))

# Bottom Outlier Note
add_card(slide4, Inches(0.8), Inches(5.9), Inches(11.7), Inches(0.9))
tb_out = slide4.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.6))
p = tb_out.text_frame.paragraphs[0]
p.text = "OUTLIER IDENTIFICADO:"
p.font.bold = True
p.font.size = Pt(12)
p.font.color.rgb = COLOR_TEXT_MAIN

p_sub = tb_out.text_frame.add_paragraph()
p_sub.text = "O caso CAD009 (Iota Autopeças) atingiu 30,5 horas de Lead Time por ter sido enviado com dados incompletos em uma sexta-feira à tarde."
p_sub.font.size = Pt(12)
p_sub.font.color.rgb = COLOR_TEXT_MUTED

# ==========================================
# SLIDE 5: ROOT CAUSE (5 Whys Card Stack)
# ==========================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide5, COLOR_BG)
add_header(slide5, "Análise de Causa Raiz (Metodologia dos 5 Porquês)")

whys = [
    ("1. Qual é o problema central?", "Lead Time total elevado (3,32h em média) e alta variação no atendimento."),
    ("2. Por que o tempo é alto?", "Porque 93% do tempo total é gasto com solicitações paradas na fila de espera."),
    ("3. Por que a fila acumula?",
     "Porque a equipe gasta horas gerenciando trocas de e-mails para corrigir dados incorretos."),
    ("4. Por que há dados incorretos?",
     "Porque o formato de entrada (Ficha Word via E-mail) aceita formulários em branco/incompletos."),
    ("CAUSA RAIZ IDENTIFICADA",
     "Ausência de uma porta de entrada digital com travas automáticas e validações em tempo real na origem.")
]

y_start = Inches(1.5)
card_h5 = Inches(0.95)

for i, (q, a) in enumerate(whys):
    is_root = (i == 4)
    bg = COLOR_DARK_BG if is_root else COLOR_CARD_BG
    txt_col = RGBColor(255, 255, 255) if is_root else COLOR_TEXT_MAIN
    q_col = COLOR_ACCENT if not is_root else RGBColor(56, 189, 248)

    y = y_start + i * Inches(1.1)
    add_card(slide5, Inches(0.8), y, Inches(11.7), card_h5, bg_color=bg)

    tb = slide5.shapes.add_textbox(Inches(1.1), y + Inches(0.12), Inches(11.1), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = q.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = q_col

    p1 = tf.add_paragraph()
    p1.text = a
    p1.font.size = Pt(13)
    p1.font.bold = is_root
    p1.font.color.rgb = txt_col

# ==========================================
# SLIDE 6: TO-BE PLAN (3 Pillar Cards)
# ==========================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide6, COLOR_BG)
add_header(slide6, "Plano de Ação Proposto (Estrutura To-Be)")

pillars = [
    ("Quick Win (Curto Prazo)", "Formulário Digital Inteligente",
     "Substituir a ficha em Word por formulário online (Forms/Typeform).\n\n"
     "• Validação de CNPJ e CEP em tempo real.\n"
     "• Campos obrigatórios (elimina 100% de fichas incompletas).\n"
     "• Anexo obrigatório de documentos."),

    ("Automação (Médio Prazo)", "Integração via API / RPA",
     "Robô de triagem e pré-cadastro no TOTVS Datasul.\n\n"
     "• Leitura direta dos dados do formulário.\n"
     "• Criação automática do rascunho de cadastro.\n"
     "• Analista faz apenas a aprovação final."),

    ("Governança (Contínuo)", "SLA & Visibilidade",
     "Gestão à vista da fila e regras claras de SLA.\n\n"
     "• Dashboard no Power BI com fila em tempo real.\n"
     "• SLA meta: < 1 hora para chamados sem pendência.\n"
     "• Notificação automática ao solicitante.")
]

col_w = Inches(3.7)
col_h = Inches(5.0)

for i, (badge, title, body) in enumerate(pillars):
    x = start_x + i * (col_w + Inches(0.3))
    add_card(slide6, x, Inches(1.6), col_w, col_h)

    tb = slide6.shapes.add_textbox(x + Inches(0.25), Inches(1.8), col_w - Inches(0.5), col_h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p_b = tf.paragraphs[0]
    p_b.text = badge.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_ACCENT
    p_b.space_after = Pt(8)

    p_t = tf.add_paragraph()
    p_t.text = title
    p_t.font.size = Pt(18)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_MAIN
    p_t.space_after = Pt(16)

    p_d = tf.add_paragraph()
    p_d.text = body
    p_d.font.size = Pt(12)
    p_d.font.color.rgb = COLOR_TEXT_MUTED

# ==========================================
# SLIDE 7: ROI & IMPACT (2 Large Impact Cards)
# ==========================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide7, COLOR_BG)
add_header(slide7, "Projeção de Impacto e Retorno (ROI)")

# Left Card - Lead Time Impact
add_card(slide7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
tb = slide7.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.1), Inches(4.4))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "REDUÇÃO DO LEAD TIME"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

p_num = tf.add_paragraph()
p_num.text = "-55%"
p_num.font.size = Pt(54)
p_num.font.bold = True
p_num.font.color.rgb = COLOR_ACCENT
p_num.space_after = Pt(10)

p_d = tf.add_paragraph()
p_d.text = "De 3,32 horas para ~1,5 hora de Lead Time médio geral.\n\n" \
           "Ao eliminar a taxa de incorreções na origem, 97% das solicitações fluem direto para a fila de execução sem paradas por troca de e-mails."
p_d.font.size = Pt(13)
p_d.font.color.rgb = COLOR_TEXT_MUTED

# Right Card - Quality & Capacity
add_card(slide7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
tb2 = slide7.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.4))
tf2 = tb2.text_frame
tf2.word_wrap = True

p = tf2.paragraphs[0]
p.text = "QUEDA NO RETRABALHO"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = RGBColor(16, 185, 129)

p_num2 = tf2.add_paragraph()
p_num2.text = "< 3%"
p_num2.font.size = Pt(54)
p_num2.font.bold = True
p_num2.font.color.rgb = RGBColor(16, 185, 129)
p_num2.space_after = Pt(10)

p_d2 = tf2.add_paragraph()
p_d2.text = "Redução drástica do retrabalho (de 30% para menos de 3%).\n\n" \
            "Aumento da capacidade produtiva da equipe atual de analistas sem necessidade de contratação de novos profissionais."
p_d2.font.size = Pt(13)
p_d2.font.color.rgb = COLOR_TEXT_MUTED

# ==========================================
# SLIDE 8: ROADMAP (Horizontal Timeline)
# ==========================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide8, COLOR_BG)
add_header(slide8, "Cronograma de Implementação")

roadmap = [
    ("Fase 1 (Semanas 1 e 2)", "Desenvolvimento & Trava",
     "Mapeamento dos campos obrigatórios e criação do formulário digital com travas de validação."),
    ("Fase 2 (Semanas 3 e 4)", "Go-Live & Treinamento",
     "Treinamento das equipes comerciais e desativação do recebimento via ficha Word."),
    ("Fase 3 (Mês 2)", "Governança & RPA",
     "Implantação do Dashboard Power BI e desenvolvimento do robô para criação automática no TOTVS.")
]

box_w = Inches(3.7)
box_h = Inches(4.5)

for i, (fase, title, desc) in enumerate(roadmap):
    x = start_x + i * (box_w + Inches(0.3))
    add_card(slide8, x, Inches(1.8), box_w, box_h)

    tb = slide8.shapes.add_textbox(x + Inches(0.3), Inches(2.1), box_w - Inches(0.6), box_h - Inches(0.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p_f = tf.paragraphs[0]
    p_f.text = fase.upper()
    p_f.font.size = Pt(10)
    p_f.font.bold = True
    p_f.font.color.rgb = COLOR_ACCENT
    p_f.space_after = Pt(8)

    p_t = tf.add_paragraph()
    p_t.text = title
    p_t.font.size = Pt(18)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_MAIN
    p_t.space_after = Pt(14)

    p_d = tf.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(13)
    p_d.font.color.rgb = COLOR_TEXT_MUTED

# Save presentation
output_path = "Apresentacao_Executiva_Cadastro_Clientes.pptx"
prs.save(output_path)
print(f"Apresentação executiva criada com sucesso em: {output_path}")